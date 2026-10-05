"""
Pillar 4.1 -- Streamlit Conversational Interface & API Integration

Conversational BI interface for business stakeholders at The Gadget Store.
"""

import uuid
import json
import io
import requests
import pandas as pd
import streamlit as st
import os
import sys
from datetime import datetime, timezone

sys.path.append(os.path.dirname(__file__))
from charts import render_chart

BRAND_BLUE = "#2E6FF2"
MONEY_WORDS = ("revenue", "profit", "price", "cost", "value", "aov")
RATIO_WORDS = ("margin", "growth", "rate", "share", "ratio", "retention")
PCT_WORDS = ("pct", "percent")

FEEDBACK_LOG_PATH = os.path.join(
    os.path.dirname(__file__), "..", "..", "logs", "feedback.log"
)

# Suggested follow-ups are drawn ONLY from questions already verified in our
# benchmark/holdout testing, so a suggestion never leads to a shaky answer.
FOLLOWUPS = {
    "revenue": [
        "What was revenue by year?",
        "What was revenue by month in 2019?",
        "What was the revenue growth from 2018 to 2019?",
        "Which continent generates the most revenue?",
    ],
    "products": [
        "What are the top 5 product categories by revenue?",
        "Show me units sold by brand",
        "Which product category has the highest average price?",
        "What's the price difference between our cheapest and most expensive product?",
    ],
    "customers": [
        "What percentage of our customers are male?",
        "What's the average age of our customers?",
        "How many customers live in London?",
        "How many distinct customers have placed an order?",
    ],
    "overview": [
        "What's our average order value?",
        "How many total orders have we had?",
        "How many stores do we have?",
        "What was total revenue in 2019?",
    ],
}


def suggest_followups(question: str, asked: list, n: int = 3) -> list:
    q = question.lower()
    if any(k in q for k in ["customer", "age", "male", "gender"]):
        topic = "customers"
    elif any(k in q for k in ["product", "categor", "brand", "price", "units"]):
        topic = "products"
    elif any(k in q for k in ["revenue", "sales", "growth"]):
        topic = "revenue"
    else:
        topic = "overview"
    ordered = FOLLOWUPS[topic] + [
        s for t, lst in FOLLOWUPS.items() if t != topic for s in lst
    ]
    asked_lower = {a.lower() for a in asked}
    return [s for s in ordered if s.lower() not in asked_lower][:n]


def log_feedback(question: str, summary: str, vote: str):
    """
    Appends a thumbs-up/down vote to a local log file (logs/feedback.log).
    Fine for this stage; a deployed multi-user product would store votes in
    a database so they persist and can be aggregated.
    """
    os.makedirs(os.path.dirname(FEEDBACK_LOG_PATH), exist_ok=True)
    entry = {
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "question": question,
        "summary": summary,
        "vote": vote,
    }
    with open(FEEDBACK_LOG_PATH, "a", encoding="utf-8") as f:
        f.write(json.dumps(entry) + "\n")


def keyed_container(key: str):
    """st.container(key=...) gives us a CSS class to target (Streamlit >= 1.39)."""
    try:
        return st.container(key=key)
    except TypeError:
        return st.container()


def is_money_col(name) -> bool:
    return any(w in str(name).lower() for w in MONEY_WORDS)


def column_kind(name, series: pd.Series) -> str:
    """How a numeric column should be displayed (matches the backend formatter)."""
    n = str(name).lower()
    if "year" in n:
        return "year"
    if any(w in n for w in PCT_WORDS):
        return "percent_points"
    if any(w in n for w in RATIO_WORDS) and "exchange" not in n:
        if series.abs().max() <= 5:
            return "percent_fraction"  # 0.58 -> 58.00%
        if "margin" in n:
            return "money"
        return "percent_points"
    if is_money_col(n):
        return "money"
    return "integer" if pd.api.types.is_integer_dtype(series) else "number"


def prettify_df(df: pd.DataFrame) -> pd.DataFrame:
    """Turns numeric-looking text columns into real numbers (no rounding here)."""
    out = df.copy()
    for col in out.columns:
        converted = pd.to_numeric(out[col], errors="coerce")
        if converted.notna().all():
            out[col] = (
                converted.astype("int64") if (converted % 1 == 0).all() else converted
            )
    return out


def format_display_df(df: pd.DataFrame) -> pd.DataFrame:
    """Display-only formatting: $ for money, % for ratios, 2 decimals, thousands separators."""
    out = prettify_df(df)
    for col in out.columns:
        if not pd.api.types.is_numeric_dtype(out[col]):
            continue
        kind = column_kind(col, out[col])
        if kind == "year":
            out[col] = out[col].astype(str)
        elif kind == "money":
            out[col] = out[col].map(lambda v: f"${v:,.2f}")
        elif kind == "percent_fraction":
            out[col] = out[col].map(lambda v: f"{v * 100:,.2f}%")
        elif kind == "percent_points":
            out[col] = out[col].map(lambda v: f"{v:,.2f}%")
        elif kind == "integer":
            out[col] = out[col].map(lambda v: f"{v:,}")
        else:
            out[col] = out[col].map(lambda v: f"{v:,.2f}")
    return out


CELL_FORMATS = {
    "year": "0",
    "money": '"$"#,##0.00',
    "percent_fraction": "0.00%",
    "percent_points": '0.00"%"',
    "integer": "#,##0",
    "number": "#,##0.00",
}
AXIS_FORMATS = {
    "money": '"$"#,##0',
    "percent_fraction": "0%",
    "percent_points": '0"%"',
}


def _stored_value(kind: str, value):
    """Keeps cell values tidy: 2 decimals (4 for fractions that display as %)."""
    if isinstance(value, float):
        if kind == "percent_fraction":
            return round(value, 4)
        if kind in ("money", "number", "percent_points"):
            return round(value, 2)
    return value


def build_excel_report(messages: list) -> bytes:
    """
    Builds ONE formatted Excel sheet covering EVERY question answered in the
    conversation. Each question gets its own section: heading, question,
    answer, formatted table and a native (editable) Excel chart in the brand
    blue. No image export involved, so it builds instantly.
    """
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.pagebreak import Break
    from openpyxl.worksheet.properties import PageSetupProperties

    # 1. Pair each question with its answer. A clarification reply is merged
    #    into the question it clarified, e.g. "What were our sales last month? (revenue)".
    items = []
    pending_question = None
    prev_status = None
    for msg in messages:
        if msg["role"] == "user":
            if prev_status == "needs_clarification" and pending_question:
                pending_question = f"{pending_question} ({msg['content']})"
            else:
                pending_question = msg["content"]
            prev_status = None
        else:
            resp = msg.get("response_data", {})
            prev_status = resp.get("status")
            if prev_status == "ok" and pending_question:
                items.append((pending_question, resp))
                pending_question = None

    BAND_FILL = PatternFill(start_color="2E6FF2", end_color="2E6FF2", fill_type="solid")
    BAND_FONT = Font(color="FFFFFF", bold=True, size=12)
    TITLE_FONT = Font(bold=True, size=16, color="12173B")
    LABEL_FONT = Font(bold=True, size=11, color="12173B")
    NCOLS = 7

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.sheet_view.showGridLines = False
    for i in range(1, NCOLS + 1):
        ws.column_dimensions[get_column_letter(i)].width = 24

    ws["A1"] = "Talk-to-my-Data -- Report"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (
        f"{len(items)} question(s) answered - generated {datetime.now().strftime('%d %b %Y, %H:%M')}"
        if items
        else "No completed answers yet to export."
    )

    row = 4
    for n, (question, resp) in enumerate(items, start=1):
        if n > 1:
            ws.row_breaks.append(
                Break(id=row - 1)
            )  # each question starts a new printed page
        # Section heading band
        for c in range(1, NCOLS + 1):
            ws.cell(row=row, column=c).fill = BAND_FILL
        ws.cell(row=row, column=1, value=f"Question {n}").font = BAND_FONT
        row += 1

        # Question and answer text
        for label, text in (
            ("Question", question),
            ("Answer", resp.get("summary", "")),
        ):
            ws.cell(row=row, column=1, value=label).font = LABEL_FONT
            ws.cell(row=row, column=1).alignment = Alignment(vertical="top")
            ws.cell(row=row, column=2, value=text)
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=NCOLS)
            ws.cell(row=row, column=2).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
            ws.row_dimensions[row].height = 34 if len(str(text)) > 110 else 20
            row += 1
        row += 1

        # Table (a single value is already stated in the answer line)
        table = resp.get("table", [])
        single_value = len(table) == 1 and len(table[0]) == 1
        if table and not single_value:
            df = prettify_df(pd.DataFrame(table))
            cols = list(df.columns)
            kinds = {
                c: column_kind(c, df[c])
                for c in cols
                if pd.api.types.is_numeric_dtype(df[c])
            }

            header_row = row
            for col_idx, col_name in enumerate(cols, start=1):
                cell = ws.cell(row=header_row, column=col_idx, value=str(col_name))
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = PatternFill(
                    start_color="12173B", end_color="12173B", fill_type="solid"
                )

            for r_off, record in enumerate(df.itertuples(index=False), start=1):
                for c_idx, value in enumerate(record, start=1):
                    value = value.item() if hasattr(value, "item") else value
                    kind = kinds.get(cols[c_idx - 1])
                    cell = ws.cell(
                        row=header_row + r_off,
                        column=c_idx,
                        value=_stored_value(kind, value) if kind else value,
                    )
                    if kind:
                        cell.number_format = CELL_FORMATS[kind]

            last_row = header_row + len(df)
            row = last_row + 2

            chart_config = resp.get("chart")
            if chart_config:
                try:
                    x_field = chart_config.get("x_field")
                    y_field = chart_config.get("y_field")
                    if x_field in cols and y_field in cols:
                        is_line = chart_config.get("type") == "line"
                        xl_chart = LineChart() if is_line else BarChart()
                        xl_chart.add_data(
                            Reference(
                                ws,
                                min_col=cols.index(y_field) + 1,
                                min_row=header_row,
                                max_row=last_row,
                            ),
                            titles_from_data=True,
                        )
                        xl_chart.set_categories(
                            Reference(
                                ws,
                                min_col=cols.index(x_field) + 1,
                                min_row=header_row + 1,
                                max_row=last_row,
                            )
                        )
                        xl_chart.legend = None
                        xl_chart.height = 7.5
                        xl_chart.width = 18
                        xl_chart.x_axis.delete = False
                        xl_chart.y_axis.delete = False

                        series = xl_chart.series[0]
                        series.graphicalProperties.solidFill = "2E6FF2"
                        series.graphicalProperties.line.solidFill = "2E6FF2"
                        if is_line:
                            series.marker.symbol = "circle"
                            series.marker.size = 7
                            series.smooth = False
                        axis_fmt = AXIS_FORMATS.get(kinds.get(y_field))
                        if axis_fmt:
                            xl_chart.y_axis.number_format = axis_fmt

                        ws.add_chart(xl_chart, f"A{row}")
                        row += 17  # room for the chart before the next question
                except Exception:
                    # A chart problem should never block the report itself
                    pass

        row += 1  # breathing room between questions

    # Print-friendly: landscape, one page wide
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# -----------------------------------------------------------------------------
# Page configuration & styling
# -----------------------------------------------------------------------------
st.set_page_config(
    page_title="Talk-to-my-Data (GenAI BI)",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
<style>
@import url('https://fonts.googleapis.com/css2?family=Source+Serif+4:wght@600;700&family=Inter:wght@400;500;600&display=swap');
 
.stApp { background-color: #F4F7FF; }
 
/* Logo alignment knob: main content top padding. Sidebar content has 2.1rem.
   If the two TGS marks are still off, nudge 5.1rem up or down in small steps. */
.block-container { padding-top: 5.1rem; }
section[data-testid="stSidebar"] > div { padding-top: 2.1rem; }
 
section[data-testid="stSidebar"] { background-color: #1C2B5C; }
section[data-testid="stSidebar"] * { color: #CADCFC !important; }
section[data-testid="stSidebar"] .stButton button,
section[data-testid="stSidebar"] .stDownloadButton button {
    background-color: #24366E;
    color: #FFFFFF !important;
    border: 1px solid #35478C;
    border-radius: 8px;
}
section[data-testid="stSidebar"] .stButton button:hover,
section[data-testid="stSidebar"] .stDownloadButton button:hover {
    background-color: #2E6FF2;
    border-color: #2E6FF2;
}
 
/* Sidebar: wrapper fills the height, bottom group is pushed to the bottom */
.st-key-sidebar_wrap { min-height: calc(100vh - 7rem); }
.st-key-sidebar_bottom,
div:has(> .st-key-sidebar_bottom) { margin-top: auto; }
 
/* Sidebar content */
.side-label {
    font-family: 'Inter', sans-serif;
    font-size: 0.8rem;
    margin-bottom: 0.1rem;
}
.snap-grid { display: grid; grid-template-columns: 1fr 1fr; gap: 8px; }
.snap {
    background-color: #24366E;
    border-radius: 10px;
    padding: 0.6rem 0.8rem;
}
section[data-testid="stSidebar"] .snap-num {
    font-family: 'Inter', sans-serif;
    font-size: 1.05rem;
    font-weight: 600;
    color: #FFFFFF !important;
}
section[data-testid="stSidebar"] .snap-lbl {
    font-family: 'Inter', sans-serif;
    font-size: 0.72rem;
    color: #A7B3DE !important;
}
.st-key-recent_list button {
    font-size: 0.82rem;
    justify-content: flex-start;
    text-align: left;
}
 
.main-title {
    font-family: 'Source Serif 4', Georgia, serif;
    font-size: 2.6rem;
    font-weight: 700;
    color: #12173B;
    letter-spacing: -1px;
}
.sub-title {
    font-family: 'Inter', sans-serif;
    font-size: 1.1rem;
    color: #6B7290;
    margin-bottom: 1.8rem;
}
 
.summary-box {
    background-color: #FFFFFF;
    border-left: 4px solid #2E6FF2;
    padding: 1.1rem 1.4rem;
    border-radius: 10px;
    font-family: 'Inter', sans-serif;
    color: #12173B;
    margin-bottom: 1rem;
    box-shadow: 0 2px 10px rgba(18,23,59,0.06);
}
.answer-label {
    font-size: 0.8rem;
    font-weight: 600;
    color: #2E6FF2;
    margin-bottom: 0.2rem;
}
.answer-text { font-size: 1.05rem; line-height: 1.6; }
.answer-text.big { font-size: 1.7rem; font-weight: 600; line-height: 1.3; }
 
.clarification-box {
    background-color: #FFFFFF;
    border-left: 4px solid #2DD4BF;
    padding: 1.2rem 1.4rem;
    border-radius: 10px;
    font-family: 'Inter', sans-serif;
    color: #12173B;
    margin-bottom: 1rem;
    box-shadow: 0 2px 10px rgba(18,23,59,0.06);
}
 
div[data-testid="stDataFrame"] {
    border-radius: 10px;
    overflow: hidden;
    box-shadow: 0 2px 10px rgba(18,23,59,0.06);
}
 
.welcome-card {
    background-color: #FFFFFF;
    border-radius: 12px;
    padding: 1.3rem;
    box-shadow: 0 2px 10px rgba(18,23,59,0.06);
    height: 100px;
    display: flex;
    flex-direction: column;
    justify-content: space-between;
}
 
[class*="st-key-feedback_"] {
    background-color: #EAF1FF;
    border-radius: 10px;
    padding: 0.7rem 1rem;
    margin-top: 0.5rem;
}
</style>
""",
    unsafe_allow_html=True,
)


# -----------------------------------------------------------------------------
# "About this data" popup
# -----------------------------------------------------------------------------
@st.dialog("About this data")
def show_about_data():
    st.markdown("##### Five connected tables")
    c1, c2 = st.columns(2)
    with c1:
        st.markdown("**Sales**")
        st.caption("Every transaction")
        st.markdown("**Products**")
        st.caption("Catalog, prices, categories")
        st.markdown("**Customers**")
        st.caption("Location, demographics")
    with c2:
        st.markdown("**Stores**")
        st.caption("66 locations + online")
        st.markdown("**Exchange Rates**")
        st.caption("Converts everything to USD")

    st.divider()
    st.markdown("##### Core KPIs tracked")
    st.markdown(
        "- Revenue over time\n- Top products & categories\n- Average Order Value\n- Category sales share"
    )

    st.divider()
    st.caption("Data covers 2016–2021. All figures are in USD.")


# -----------------------------------------------------------------------------
# Session state
# -----------------------------------------------------------------------------
if "conversation_id" not in st.session_state:
    st.session_state.conversation_id = str(uuid.uuid4())[:8]
if "messages" not in st.session_state:
    st.session_state.messages = []
if "pending_clarification" not in st.session_state:
    st.session_state.pending_clarification = False
if "last_question" not in st.session_state:
    st.session_state.last_question = ""
if "feedback_given" not in st.session_state:
    st.session_state.feedback_given = {}

backend_url = os.getenv("BACKEND_API_URL", "http://127.0.0.1:8080/ask")


# -----------------------------------------------------------------------------
# Sidebar
# -----------------------------------------------------------------------------
with st.sidebar:
    with keyed_container("sidebar_wrap"):
        st.markdown(
            """
<div style="font-family: 'Source Serif 4', Georgia, serif; font-size: 1.6rem; font-weight: 700; color: #FFFFFF;">
    T<span style="color:#2E6FF2;">G</span>S
</div>
<div style="font-family: 'Inter', sans-serif; font-size: 0.85rem; color: #A7B3DE;">
    GenAI Business Intelligence Assistant
</div>
""",
            unsafe_allow_html=True,
        )

        st.divider()

        if st.button("ℹ️ About this data", use_container_width=True):
            show_about_data()

        # Excel report: native chart, no image export, so it builds instantly
        if st.session_state.messages:
            try:
                st.download_button(
                    "⬇️ Download Excel report",
                    data=build_excel_report(st.session_state.messages),
                    file_name="talk_to_my_data_report.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                )
            except Exception:
                st.caption("Excel report unavailable for this conversation.")

        # Data snapshot -- static facts verified in our EDA
        st.markdown(
            '<div class="side-label">Data snapshot</div>', unsafe_allow_html=True
        )
        st.markdown(
            """
<div class="snap-grid">
    <div class="snap"><div class="snap-num">2016–21</div><div class="snap-lbl">Data range</div></div>
    <div class="snap"><div class="snap-num">62,884</div><div class="snap-lbl">Transactions</div></div>
    <div class="snap"><div class="snap-num">15,266</div><div class="snap-lbl">Customers</div></div>
    <div class="snap"><div class="snap-num">9</div><div class="snap-lbl">Markets</div></div>
</div>
""",
            unsafe_allow_html=True,
        )

        # Recent questions -- click to ask again
        asked_questions = [
            m["content"] for m in st.session_state.messages if m["role"] == "user"
        ]
        if asked_questions:
            st.markdown(
                '<div class="side-label">Recent questions</div>', unsafe_allow_html=True
            )
            recent = []
            for q in reversed(asked_questions):
                if q not in recent:
                    recent.append(q)
                if len(recent) == 5:
                    break
            with keyed_container("recent_list"):
                for i, q in enumerate(recent):
                    label = q if len(q) <= 55 else q[:52] + "…"
                    if st.button(
                        label, key=f"recent_{i}", help=q, use_container_width=True
                    ):
                        st.session_state.queued_question = q

        # Bottom group: pushed to the bottom of the sidebar
        with keyed_container("sidebar_bottom"):
            st.divider()
            if st.button("🗑️ Clear Conversation", use_container_width=True):
                st.session_state.messages = []
                st.session_state.conversation_id = str(uuid.uuid4())[:8]
                st.session_state.pending_clarification = False
                st.session_state.last_question = ""
                st.session_state.feedback_given = {}
                st.rerun()

            with st.expander("⚙️ Connection settings"):
                backend_url = st.text_input("Backend API URL", value=backend_url)
                health_url = backend_url.rsplit("/ask", 1)[0] + "/health"
                try:
                    health_resp = requests.get(health_url, timeout=3)
                    if health_resp.status_code == 200:
                        st.success("Connected")
                    else:
                        st.warning("Degraded")
                except Exception:
                    st.error("Offline -- start the backend")


# -----------------------------------------------------------------------------
# Main header
# -----------------------------------------------------------------------------
st.markdown(
    """
<div style="margin-bottom: 0.3rem;">
    <div class="main-title">T<span style="color:#2E6FF2;">G</span>S</div>
</div>
<div class="sub-title">Faster decisions. Better results. Ask a question, in plain English.</div>
""",
    unsafe_allow_html=True,
)

# A question picked from "Recent questions" in the sidebar, if any
selected_sample = st.session_state.pop("queued_question", None)

# -----------------------------------------------------------------------------
# Welcome state
# -----------------------------------------------------------------------------
if not st.session_state.messages:
    welcome_questions = [
        ("📈", "What was total revenue in December 2019?"),
        ("🏆", "What are the top 5 product categories by revenue?"),
        ("💰", "What's our average order value?"),
    ]
    cols = st.columns(3)
    for col, (icon, q) in zip(cols, welcome_questions):
        with col:
            st.markdown(
                f"""
<div class="welcome-card">
    <div style="font-size: 1.4rem;">{icon}</div>
    <div style="font-family: 'Inter', sans-serif; font-size: 0.92rem; color: #12173B; font-weight: 500;">
        {q}
    </div>
</div>
""",
                unsafe_allow_html=True,
            )
            if st.button("Ask this", key=f"welcome_{q}", use_container_width=True):
                selected_sample = q
    st.markdown("<div style='margin-bottom:1rem;'></div>", unsafe_allow_html=True)


# -----------------------------------------------------------------------------
# Chat history
# -----------------------------------------------------------------------------
for idx, msg in enumerate(st.session_state.messages):
    role = msg["role"]
    with st.chat_message(role):
        if role == "user":
            st.write(msg["content"])
        else:
            resp = msg.get("response_data", {})
            status = resp.get("status", "error")

            if status == "ok":
                question_text = (
                    st.session_state.messages[idx - 1]["content"] if idx > 0 else ""
                )
                table_data = resp.get("table", [])
                # A single value is already stated in the answer -- showing it
                # again as a raw table is just noise.
                single_value = len(table_data) == 1 and len(table_data[0]) == 1

                if resp.get("summary"):
                    big = " big" if single_value else ""
                    st.markdown(
                        f'<div class="summary-box"><div class="answer-label">Answer</div>'
                        f'<div class="answer-text{big}">{resp["summary"]}</div></div>',
                        unsafe_allow_html=True,
                    )

                if table_data and not single_value:
                    st.dataframe(
                        format_display_df(pd.DataFrame(table_data)),
                        use_container_width=True,
                        hide_index=True,
                    )

                chart_config = resp.get("chart")
                if chart_config and table_data:
                    render_chart(chart_config, table_data)

                if resp.get("sql"):
                    with st.expander("🔍 View Generated SQL Query"):
                        st.code(resp["sql"], language="sql")

                # Feedback, in its own labeled box
                with keyed_container(f"feedback_{idx}"):
                    if idx not in st.session_state.feedback_given:
                        st.caption("Was this answer correct?")
                        fcol1, fcol2, _ = st.columns([1, 1, 6])
                        with fcol1:
                            if st.button("👍 Yes", key=f"up_{idx}"):
                                log_feedback(
                                    question_text, resp.get("summary", ""), "up"
                                )
                                st.session_state.feedback_given[idx] = "up"
                                st.rerun()
                        with fcol2:
                            if st.button("👎 No", key=f"down_{idx}"):
                                log_feedback(
                                    question_text, resp.get("summary", ""), "down"
                                )
                                st.session_state.feedback_given[idx] = "down"
                                st.rerun()
                    else:
                        vote = st.session_state.feedback_given[idx]
                        st.caption(
                            "✅ Thanks for confirming this was correct."
                            if vote == "up"
                            else "📝 Thanks -- flagged for review."
                        )

                # Suggested follow-ups, only under the most recent answer
                if idx == len(st.session_state.messages) - 1:
                    suggestions = suggest_followups(question_text, asked_questions)
                    if suggestions:
                        st.caption("You might also ask")
                        scols = st.columns(len(suggestions))
                        for i, (scol, s) in enumerate(zip(scols, suggestions)):
                            with scol:
                                if st.button(
                                    s, key=f"fu_{idx}_{i}", use_container_width=True
                                ):
                                    selected_sample = s

            elif status == "needs_clarification":
                clarify_q = resp.get(
                    "clarification_question", "The query requires clarification."
                )
                st.markdown(
                    f'<div class="clarification-box">❓ <b>Clarification Needed:</b><br>{clarify_q}</div>',
                    unsafe_allow_html=True,
                )
                st.caption("Please type your answer below to continue.")

            elif status == "error":
                err_msg = resp.get(
                    "message", "An error occurred while processing your request."
                )
                st.error(f"⚠️ {err_msg}")


# -----------------------------------------------------------------------------
# New user input
# -----------------------------------------------------------------------------
user_input = st.chat_input(
    "Type your business question here (e.g., What were total sales in 2025?)..."
)

if selected_sample:
    user_input = selected_sample
    # A clicked suggestion is a fresh question, not an answer to a clarification
    st.session_state.pending_clarification = False

if user_input:
    st.session_state.messages.append({"role": "user", "content": user_input})
    with st.chat_message("user"):
        st.write(user_input)

    payload = {
        "question": user_input,
        "conversation_id": st.session_state.conversation_id,
        "clarification_answer": None,
    }

    if st.session_state.pending_clarification and st.session_state.last_question:
        payload["question"] = st.session_state.last_question
        payload["clarification_answer"] = user_input
        st.session_state.pending_clarification = False
    else:
        st.session_state.last_question = user_input

    with st.chat_message("assistant"):
        with st.spinner("Analyzing question, generating SQL & querying BigQuery..."):
            try:
                api_response = requests.post(
                    backend_url,
                    json=payload,
                    headers={"Content-Type": "application/json"},
                    timeout=90,
                )
                if api_response.status_code == 200:
                    data = api_response.json()
                else:
                    data = {
                        "status": "error",
                        "message": f"Server returned HTTP {api_response.status_code}: {api_response.text}",
                    }
            except requests.exceptions.Timeout:
                data = {
                    "status": "error",
                    "message": "Backend request timed out (90s limit). Please check server health or retry.",
                }
            except Exception as e:
                data = {
                    "status": "error",
                    "message": f"Could not connect to backend API: {str(e)}",
                }

        if data.get("status") == "needs_clarification":
            st.session_state.pending_clarification = True

        st.session_state.messages.append(
            {"role": "assistant", "content": "", "response_data": data}
        )
        st.rerun()
